#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Générateur de fiches "Visa de l'Expert-Comptable"
==================================================

Flux de l'application :
  1) Importer un fichier Excel (feuille avec un ou plusieurs tableaux
     "Critères / Choix" empilés, un tableau = un client) -> JSON brut.
  2) Envoyer ce JSON brut à Gemini pour le restructurer proprement en :
         {
           "client 1": { "raison_sociale": "...", "nom_pdg": "...", ... },
           "client 2": { ... }
         }
  3) Générer un fichier Word (.docx) par client à partir du gabarit
     `template.docx` (copie de votre modèle ABC COMPANY avec des
     variables Jinja {{ ... }} à la place des infos spécifiques au
     client).

Dépendances (voir requirements.txt) :
    pip install openpyxl docxtpl requests
"""

import json
import os
import re
import sys
import threading
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext

import openpyxl
import requests
from docxtpl import DocxTemplate


# --------------------------------------------------------------------------
# Chemins
# --------------------------------------------------------------------------

def resource_path(filename: str) -> str:
    """Retourne le chemin absolu d'un fichier livré à côté du script
    (fonctionne aussi une fois empaqueté avec PyInstaller)."""
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, filename)


TEMPLATE_PATH = resource_path("template.docx")
DEFAULT_OUTPUT_DIR = os.path.join(os.path.expanduser("~"), "Fiches_Visa_Generees")


# --------------------------------------------------------------------------
# 1) EXCEL -> JSON BRUT
# --------------------------------------------------------------------------

CRIT_RE = re.compile(r"crit[eè]res?", re.IGNORECASE)
CHOIX_RE = re.compile(r"choix", re.IGNORECASE)


def _norm(v) -> str:
    return "" if v is None else str(v).strip()


def _build_merged_map(ws):
    """Associe chaque cellule fusionnée à la cellule 'maître' (haut-gauche)
    qui porte réellement la valeur."""
    merged_map = {}
    for rng in ws.merged_cells.ranges:
        master = ws.cell(row=rng.min_row, column=rng.min_col)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged_map[(r, c)] = master
    return merged_map


def _cell_value(ws, merged_map, row, col):
    cell = merged_map.get((row, col))
    if cell is None:
        cell = ws.cell(row=row, column=col)
    return cell.value


def _find_header(ws, merged_map, row, max_col):
    """Si `row` est une ligne d'en-tête 'Critères | Choix', renvoie
    (colonne_critere, colonne_choix). Sinon None."""
    crit_col = None
    for col in range(1, max_col + 1):
        v = _norm(_cell_value(ws, merged_map, row, col))
        if v and CRIT_RE.search(v):
            crit_col = col
            break
    if crit_col is None:
        return None
    choix_col = None
    for col in range(crit_col + 1, max_col + 1):
        v = _norm(_cell_value(ws, merged_map, row, col))
        if v and CHOIX_RE.search(v):
            choix_col = col
            break
    if choix_col is None:
        return None
    return crit_col, choix_col


def extract_clients_from_excel(path: str) -> list:
    """Parcourt toutes les feuilles et détecte chaque tableau
    'Critères / Choix' empilé comme un client distinct."""
    wb = openpyxl.load_workbook(path, data_only=True)
    clients = []

    for ws in wb.worksheets:
        merged_map = _build_merged_map(ws)
        max_row, max_col = ws.max_row, ws.max_column
        row = 1
        while row <= max_row:
            header = _find_header(ws, merged_map, row, max_col)
            if not header:
                row += 1
                continue

            crit_col, choix_col = header
            block = {}
            r = row + 1
            while r <= max_row:
                label = _norm(_cell_value(ws, merged_map, r, crit_col))

                # Une nouvelle ligne d'en-tête = début du client suivant
                if label and CRIT_RE.search(label) and _find_header(ws, merged_map, r, max_col):
                    break

                if label == "":
                    # ligne vide : fin du tableau si on a déjà des données
                    if block:
                        break
                    r += 1
                    continue

                value = _norm(_cell_value(ws, merged_map, r, choix_col))
                block[label] = value
                r += 1

            if block:
                clients.append(block)
            row = r if r > row else row + 1

    return clients


# --------------------------------------------------------------------------
# 2) JSON BRUT -> JSON STRUCTURÉ (via Gemini)
# --------------------------------------------------------------------------

SCHEMA_DESCRIPTION = """\
Pour CHAQUE client, produis un objet avec EXACTEMENT ces clés (chaîne de
caractères, "" si l'information est absente) :

- raison_sociale      (ex-"Raison sociale")
- titre_pdg            "Monsieur" ou "Madame" (déduis-le du prénom si possible,
                        sinon mets "Monsieur")
- nom_pdg              (ex-"Nom de PDG", sans le "Monsieur"/"Madame")
- adresse               (ex-"Adresse de la société")
- exercice              (ex-"Exercice", ex: "2026")
- trimestre             (ex-"Trimestre", normalisé en "T1", "T2", "T3" ou "T4")
- type_attestation      (ex-"Deux types d'attestations", ex: "Avec retard")
- montant                (ex-"Montant", garde le format nombre ex "17.000,00")
- lieu                    (ex-"Lieu")
- date_signature          (ex-"Date de signature", format JJ/MM/AAAA)
- nom_signataire          (ex-"Nom & Prenom" de l'expert-comptable signataire)
- titre_signataire        (ex-"Signataire", ex: "Associé" ou "Expert-comptable")
"""


def build_gemini_prompt(raw_clients: list) -> str:
    return f"""Tu reçois ci-dessous une liste de clients extraits automatiquement
d'un fichier Excel. Chaque élément de la liste est un dictionnaire
"libellé du critère (français, tel quel dans Excel)" -> "valeur choisie".

{SCHEMA_DESCRIPTION}

Réponds UNIQUEMENT avec un objet JSON valide (pas de texte autour, pas de
balises markdown, pas de ```), de la forme :

{{
  "client 1": {{ "raison_sociale": "...", "titre_pdg": "...", ... }},
  "client 2": {{ ... }}
}}

Il doit y avoir exactement un "client N" par élément de la liste ci-dessous,
dans le même ordre.

Données brutes extraites du fichier Excel :
{json.dumps(raw_clients, ensure_ascii=False, indent=2)}
"""


def call_gemini(api_key: str, model: str, prompt: str) -> dict:
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={api_key}"
    payload = {
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {
            "temperature": 0.1,
            "responseMimeType": "application/json",
        },
    }
    resp = requests.post(url, json=payload, timeout=90)
    if resp.status_code != 200:
        raise RuntimeError(f"Erreur API Gemini ({resp.status_code}) : {resp.text[:500]}")

    data = resp.json()
    try:
        text = data["candidates"][0]["content"]["parts"][0]["text"]
    except (KeyError, IndexError) as exc:
        raise RuntimeError(f"Réponse Gemini inattendue : {json.dumps(data)[:500]}") from exc

    text = text.strip()
    text = re.sub(r"^```(json)?", "", text).strip()
    text = re.sub(r"```$", "", text).strip()

    try:
        return json.loads(text)
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"Gemini n'a pas renvoyé un JSON valide :\n{text[:800]}") from exc


# --------------------------------------------------------------------------
# 3) JSON STRUCTURÉ -> FICHIERS WORD (.docx)
# --------------------------------------------------------------------------

TRIMESTRE_INFO = {
    "T1": {"debut_mmdd": "01/01", "fin_mmdd": "31/03", "debut_mois": "JANVIER", "fin_txt": "31 MARS"},
    "T2": {"debut_mmdd": "01/04", "fin_mmdd": "30/06", "debut_mois": "AVRIL", "fin_txt": "30 JUIN"},
    "T3": {"debut_mmdd": "01/07", "fin_mmdd": "30/09", "debut_mois": "JUILLET", "fin_txt": "30 SEPTEMBRE"},
    "T4": {"debut_mmdd": "01/10", "fin_mmdd": "31/12", "debut_mois": "OCTOBRE", "fin_txt": "31 DÉCEMBRE"},
}


def build_docx_context(client: dict) -> dict:
    trimestre = _norm(client.get("trimestre", "T1")).upper()
    if trimestre not in TRIMESTRE_INFO:
        trimestre = "T1"
    info = TRIMESTRE_INFO[trimestre]
    exercice = _norm(client.get("exercice", ""))

    return {
        "raison_sociale": client.get("raison_sociale", ""),
        "titre_pdg": client.get("titre_pdg") or "Monsieur",
        "nom_pdg": client.get("nom_pdg", ""),
        "adresse": client.get("adresse", ""),
        "periode_titre_suite": f'{info["debut_mois"]} AU {info["fin_txt"]} {exercice}'.strip(),
        "periode_debut": f'{info["debut_mmdd"]}/{exercice}' if exercice else "",
        "periode_fin": f'{info["fin_mmdd"]}/{exercice}' if exercice else "",
        "montant": client.get("montant", ""),
        "lieu": client.get("lieu", "Casablanca"),
        "date_signature": client.get("date_signature", ""),
        "nom_signataire": client.get("nom_signataire", ""),
        "titre_signataire": client.get("titre_signataire") or "Associé",
    }


def sanitize_filename(name: str) -> str:
    name = re.sub(r'[\\/*?:"<>|]', "_", name).strip()
    return name or "client"


def generate_docs(structured_clients: dict, output_dir: str) -> list:
    if not os.path.exists(TEMPLATE_PATH):
        raise FileNotFoundError(
            f"Gabarit introuvable : {TEMPLATE_PATH}\n"
            "Placez 'template.docx' à côté de app.py."
        )
    os.makedirs(output_dir, exist_ok=True)
    generated = []

    for client_key, client_data in structured_clients.items():
        ctx = build_docx_context(client_data)
        doc = DocxTemplate(TEMPLATE_PATH)
        doc.render(ctx)
        filename = f"Visa_{sanitize_filename(ctx['raison_sociale'] or client_key)}.docx"
        out_path = os.path.join(output_dir, filename)
        doc.save(out_path)
        generated.append(out_path)

    return generated


def open_folder(path: str):
    try:
        if sys.platform.startswith("win"):
            os.startfile(path)  # type: ignore[attr-defined]
        elif sys.platform == "darwin":
            os.system(f'open "{path}"')
        else:
            os.system(f'xdg-open "{path}"')
    except Exception:
        pass


# --------------------------------------------------------------------------
# INTERFACE GRAPHIQUE (Tkinter)
# --------------------------------------------------------------------------

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Générateur de fiches Visa Expert-Comptable")
        self.geometry("980x800")
        self.minsize(860, 640)

        self.excel_path = None
        self.raw_clients = []
        self.structured_clients = {}
        self.output_dir = DEFAULT_OUTPUT_DIR

        self._build_ui()

    # ---------------------------------------------------------- UI layout
    def _build_ui(self):
        pad = {"padx": 10, "pady": 6}

        # ---- Zone réglages Gemini ------------------------------------
        settings = ttk.LabelFrame(self, text="Réglages Gemini")
        settings.pack(fill="x", **pad)

        ttk.Label(settings, text="Clé API Gemini :").grid(row=0, column=0, sticky="w", padx=6, pady=6)
        self.api_key_var = tk.StringVar(value=os.environ.get("GEMINI_API_KEY", ""))
        ttk.Entry(settings, textvariable=self.api_key_var, show="•", width=50).grid(
            row=0, column=1, sticky="we", padx=6, pady=6
        )

        ttk.Label(settings, text="Modèle :").grid(row=0, column=2, sticky="w", padx=6, pady=6)
        self.model_var = tk.StringVar(value="gemini-2.0-flash")
        ttk.Combobox(
            settings, textvariable=self.model_var, width=22,
            values=["gemini-2.0-flash", "gemini-2.0-flash-lite", "gemini-1.5-flash", "gemini-1.5-pro", "gemini-3.6-flash","gemini-3.5-flash","gemini-3.7-flash"],
        ).grid(row=0, column=3, sticky="w", padx=6, pady=6)
        settings.columnconfigure(1, weight=1)

        # ---- Étape 1 : Import Excel -----------------------------------
        step1 = ttk.LabelFrame(self, text="Étape 1 — Importer le fichier Excel")
        step1.pack(fill="both", expand=True, **pad)

        row1 = ttk.Frame(step1)
        row1.pack(fill="x", padx=6, pady=6)
        ttk.Button(row1, text="📂 Importer un fichier Excel", command=self.on_import_excel).pack(side="left")
        self.excel_label = ttk.Label(row1, text="Aucun fichier importé")
        self.excel_label.pack(side="left", padx=10)

        self.raw_text = scrolledtext.ScrolledText(step1, height=10, wrap="word")
        self.raw_text.pack(fill="both", expand=True, padx=6, pady=6)

        # ---- Étape 2 : Structuration Gemini -----------------------------
        step2 = ttk.LabelFrame(self, text="Étape 2 — Structurer avec Gemini")
        step2.pack(fill="both", expand=True, **pad)

        row2 = ttk.Frame(step2)
        row2.pack(fill="x", padx=6, pady=6)
        self.btn_gemini = ttk.Button(row2, text="🤖 Structurer avec Gemini", command=self.on_structure_gemini)
        self.btn_gemini.pack(side="left")
        self.gemini_status = ttk.Label(row2, text="")
        self.gemini_status.pack(side="left", padx=10)

        self.structured_text = scrolledtext.ScrolledText(step2, height=12, wrap="word")
        self.structured_text.pack(fill="both", expand=True, padx=6, pady=6)
        ttk.Label(
            step2,
            text="Vous pouvez corriger le JSON ci-dessus à la main avant de générer les fiches.",
            foreground="#555",
        ).pack(anchor="w", padx=6)

        # ---- Étape 3 : Génération Word ----------------------------------
        step3 = ttk.LabelFrame(self, text="Étape 3 — Générer les fiches Word")
        step3.pack(fill="x", **pad)

        row3 = ttk.Frame(step3)
        row3.pack(fill="x", padx=6, pady=6)
        ttk.Button(row3, text="📁 Dossier de sortie...", command=self.on_choose_output_dir).pack(side="left")
        self.out_dir_label = ttk.Label(row3, text=self.output_dir)
        self.out_dir_label.pack(side="left", padx=10)

        row3b = ttk.Frame(step3)
        row3b.pack(fill="x", padx=6, pady=6)
        ttk.Button(row3b, text="📝 Générer les fiches Word", command=self.on_generate_docs).pack(side="left")

        # ---- Barre de statut ---------------------------------------------
        self.status_var = tk.StringVar(value="Prêt.")
        ttk.Label(self, textvariable=self.status_var, relief="sunken", anchor="w").pack(fill="x", side="bottom")

    # ---------------------------------------------------------- Étape 1
    def on_import_excel(self):
        path = filedialog.askopenfilename(
            title="Choisir un fichier Excel",
            filetypes=[("Fichiers Excel", "*.xlsx *.xlsm"), ("Tous les fichiers", "*.*")],
        )
        if not path:
            return
        try:
            clients = extract_clients_from_excel(path)
        except Exception as exc:
            messagebox.showerror("Erreur d'import", str(exc))
            return

        if not clients:
            messagebox.showwarning(
                "Aucun client détecté",
                "Aucun tableau 'Critères / Choix' n'a été détecté dans ce fichier.",
            )
            return

        self.excel_path = path
        self.raw_clients = clients
        self.excel_label.config(text=f"{os.path.basename(path)}  ({len(clients)} client(s) détecté(s))")
        self.raw_text.delete("1.0", tk.END)
        self.raw_text.insert(tk.END, json.dumps(clients, ensure_ascii=False, indent=2))
        self.status_var.set(f"{len(clients)} client(s) importé(s) depuis Excel.")

    # ---------------------------------------------------------- Étape 2
    def on_structure_gemini(self):
        raw_json_text = self.raw_text.get("1.0", tk.END).strip()
        if not raw_json_text:
            messagebox.showwarning("JSON manquant", "Importez d'abord un fichier Excel (étape 1).")
            return
        try:
            raw_clients = json.loads(raw_json_text)
        except json.JSONDecodeError as exc:
            messagebox.showerror("JSON invalide", f"Le JSON brut de l'étape 1 est invalide :\n{exc}")
            return

        api_key = self.api_key_var.get().strip()
        if not api_key:
            messagebox.showwarning(
                "Clé API manquante",
                "Renseignez votre clé API Gemini en haut de la fenêtre "
                "(ou définissez la variable d'environnement GEMINI_API_KEY).",
            )
            return

        self.btn_gemini.config(state="disabled")
        self.gemini_status.config(text="Appel de Gemini en cours...")
        self.status_var.set("Appel de Gemini en cours...")

        def worker():
            try:
                prompt = build_gemini_prompt(raw_clients)
                structured = call_gemini(api_key, self.model_var.get().strip(), prompt)
            except Exception as exc:
                self.after(0, lambda: self._on_gemini_error(exc))
                return
            self.after(0, lambda: self._on_gemini_success(structured))

        threading.Thread(target=worker, daemon=True).start()

    def _on_gemini_success(self, structured: dict):
        self.structured_clients = structured
        self.structured_text.delete("1.0", tk.END)
        self.structured_text.insert(tk.END, json.dumps(structured, ensure_ascii=False, indent=2))
        self.btn_gemini.config(state="normal")
        self.gemini_status.config(text=f"{len(structured)} client(s) structuré(s).")
        self.status_var.set("Structuration Gemini terminée.")

    def _on_gemini_error(self, exc: Exception):
        self.btn_gemini.config(state="normal")
        self.gemini_status.config(text="Échec.")
        self.status_var.set("Erreur lors de l'appel à Gemini.")
        messagebox.showerror("Erreur Gemini", str(exc))

    # ---------------------------------------------------------- Étape 3
    def on_choose_output_dir(self):
        path = filedialog.askdirectory(title="Choisir le dossier de sortie", initialdir=self.output_dir)
        if path:
            self.output_dir = path
            self.out_dir_label.config(text=self.output_dir)

    def on_generate_docs(self):
        structured_text = self.structured_text.get("1.0", tk.END).strip()
        if not structured_text:
            messagebox.showwarning(
                "Rien à générer", "Structurez d'abord les données avec Gemini (étape 2),"
                " ou collez un JSON structuré valide."
            )
            return
        try:
            structured = json.loads(structured_text)
        except json.JSONDecodeError as exc:
            messagebox.showerror("JSON invalide", f"Le JSON structuré de l'étape 2 est invalide :\n{exc}")
            return

        try:
            generated = generate_docs(structured, self.output_dir)
        except Exception as exc:
            messagebox.showerror("Erreur de génération", str(exc))
            return

        self.status_var.set(f"{len(generated)} fiche(s) Word générée(s) dans {self.output_dir}")
        if messagebox.askyesno(
            "Fiches générées",
            f"{len(generated)} fiche(s) Word ont été générées dans :\n{self.output_dir}\n\n"
            "Ouvrir le dossier ?",
        ):
            open_folder(self.output_dir)


if __name__ == "__main__":
    App().mainloop()
